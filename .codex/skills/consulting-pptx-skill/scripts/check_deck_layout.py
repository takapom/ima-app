#!/usr/bin/env python3
"""pptxのレイアウト崩れを機械的に見つける。

見つけるもの
  hidden   … 後ろに描かれた不透明な図形に覆われて、文字が見えなくなっている
  empty    … 文字も画像も無い箱（消し忘れの吹き出し・枠）
  outside  … スライドの外にはみ出している
  overflow … 箱に対して文字が多すぎて溢れている
  squash   … 画像の縦横比が元と大きく違う（引き伸ばし）

使い方: python3 scripts/check_deck_layout.py <pptx> [ページ番号...]

結果はすべて WARN 扱い（終了コードは常に 0）。overflow は文字数からの見積もりなので
実際の表示と差が出る。検出するのは「中身の無い箱」で、空のページ・空の表は対象外。
HTML は対象外（HTML は check_layout.mjs）。
"""
import math
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

EMU_IN = 914400


def walk(shapes, dx=0, dy=0, sx=1.0, sy=1.0):
    for sh in shapes:
        try:
            l = dx + sh.left * sx
            t = dy + sh.top * sy
            w = sh.width * sx
            h = sh.height * sy
        except TypeError:
            continue
        yield sh, (l, t, w, h)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            g = sh._element.find(".//" + qn("a:xfrm"))
            cx = cy = 0
            cw, ch = sh.width, sh.height
            if g is not None:
                co, ce = g.find(qn("a:chOff")), g.find(qn("a:chExt"))
                if co is not None:
                    cx, cy = int(co.get("x")), int(co.get("y"))
                if ce is not None:
                    cw = int(ce.get("cx")) or sh.width
                    ch = int(ce.get("cy")) or sh.height
            rx = (sh.width / cw) if cw else 1
            ry = (sh.height / ch) if ch else 1
            yield from walk(sh.shapes, dx + (sh.left - cx * rx) * sx,
                            dy + (sh.top - cy * ry) * sy, sx * rx, sy * ry)


def overlap(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    ix = max(0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0, min(ay + ah, by + bh) - max(ay, by))
    return ix * iy


def is_opaque(sh):
    """塗りつぶし or 画像 で下が見えなくなる図形か"""
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    try:
        f = sh.fill
        if f.type is None:
            return False
        return "SOLID" in str(f.type) or "PATTERN" in str(f.type) or "PICTURE" in str(f.type)
    except Exception:
        return False


def text_of(sh):
    tf = getattr(sh, "text_frame", None)
    return tf.text.strip() if tf is not None else ""


def font_pt(sh, default=10.0):
    tf = getattr(sh, "text_frame", None)
    if tf is None:
        return default
    for p in tf.paragraphs:
        for r in p.runs:
            if r.font.size:
                return r.font.size.pt
    return default


def check(path, pages=None):
    prs = Presentation(path)
    SW, SH = prs.slide_width, prs.slide_height
    issues = []
    for n, slide in enumerate(prs.slides, 1):
        if pages and n not in pages:
            continue
        items = list(walk(slide.shapes))
        for i, (sh, box) in enumerate(items):
            t = text_of(sh)
            l, top, w, h = box
            if w <= 0 or h <= 0:
                continue
            # 外にはみ出していないか
            if l < -w * 0.5 or top < -h * 0.5 or l + w > SW + w * 0.5 or top + h > SH * 1.02:
                issues.append((n, "outside", f"{(t or sh.name)[:24]} が枠外 ({l/EMU_IN:.1f}in,{top/EMU_IN:.1f}in)"))
            if t:
                # 上に乗っている不透明な図形に覆われていないか
                covered = 0
                for sh2, box2 in items[i + 1:]:
                    if sh2 is sh or not is_opaque(sh2):
                        continue
                    if text_of(sh2) and overlap(box, box2) < w * h * 0.9:
                        continue
                    covered = max(covered, overlap(box, box2))
                if covered > w * h * 0.85:
                    issues.append((n, "hidden", f"『{t[:26]}』が他の図形に隠れている"))
                # 文字が溢れていないか（ざっくり見積もり）
                pt = font_pt(sh)
                cpl = max(1, int((w / EMU_IN * 72) / (pt * 1.02)))
                lines = sum(max(1, math.ceil(len(p.text) / cpl)) for p in sh.text_frame.paragraphs)
                need = lines * pt * 1.45 / 72 * EMU_IN
                if top + need > SH:
                    issues.append((n, "outside", f"『{t[:22]}』が溢れてスライドの外に出ている"))
                elif need > h * 1.3 and h > 100000:
                    issues.append((n, "overflow", f"『{t[:22]}』が箱から溢れている（必要{need/EMU_IN:.1f}in / 箱{h/EMU_IN:.1f}in）"))
            else:
                # 文字も画像も無い箱（枠だけの強調は許容するので、塗りがあるものだけ）
                # 高さか幅が 0.05in 未満の図形は線（罫線・区切り）なので箱として扱わない
                thin = min(w, h) < 0.05 * EMU_IN
                if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and is_opaque(sh) and not thin:
                    # 画像の上の強調枠、または上に文字が載っている下地なら問題なし
                    on_pic = any(sh2.shape_type == MSO_SHAPE_TYPE.PICTURE and overlap(box, box2) > w * h * 0.5
                                 for sh2, box2 in items)
                    has_label = any(text_of(sh2) and overlap(box, box2) > w * h * 0.3
                                    for sh2, box2 in items[i + 1:])
                    if not on_pic and not has_label:
                        issues.append((n, "empty", f"中身の無い箱 ({w/EMU_IN:.1f}x{h/EMU_IN:.1f}in)"))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                # 画像が塗りつぶしの図形に覆われていないか
                for sh2, box2 in items[i + 1:]:
                    if sh2.shape_type == MSO_SHAPE_TYPE.PICTURE or not is_opaque(sh2):
                        continue
                    if text_of(sh2):
                        continue
                    if overlap(box, box2) > w * h * 0.15:
                        issues.append((n, "hidden", f"画像の{overlap(box, box2)/(w*h)*100:.0f}%が塗りつぶしの箱に覆われている"))
                        break
                try:
                    iw, ih = sh.image.size
                    if iw and ih:
                        r = (w / h) / (iw / ih)
                        if r > 1.25 or r < 0.8:
                            issues.append((n, "squash", f"画像の縦横比が{r:.2f}倍ずれている"))
                except Exception:
                    pass
    return issues


if __name__ == "__main__":
    path = sys.argv[1]
    pages = [int(x) for x in sys.argv[2:]] or None
    issues = check(path, pages)
    for n, kind, msg in issues:
        print(f"p{n:<3} [{kind}] {msg}")
    print(f"--- {len(issues)} 件")
