#!/usr/bin/env python3
"""HTMLデッキ（16:9・1 section = 1スライド）を、編集できる PPTX に変換する。

使い方:
  python3 html_to_pptx.py deck.html                 # deck.pptx を同じフォルダに出力
  python3 html_to_pptx.py deck.html out.pptx

前提: 資料は HTML で仕上げてから変換する（slide-rules §8.6）。変換は、ユーザーが PPTX を明示的に求めたときだけ行う。
仕組み: html_dump.mjs が Chrome で描画した要素の位置と見た目を書き出し、本スクリプトが PowerPoint の図形に組み立てる。
  文字 → テキストボックス（折り返し幅・行間・書体・色を保持）
  塗り・枠 → 四角形（角丸・多角形の clip-path は形状として再現）／片側だけの罫線 → 直線
  表 → PowerPoint の表（結合セル・セルの塗り・罫線・余白を保持）
  SVG・画像・背景画像 → 透過PNG（中の数値は編集できない）
依存: Node 22+ と Chrome（html_dump.mjs）、python-pptx。
"""
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLIDE_W, SLIDE_H = 12192000, 6858000
SANS, SERIF = "Yu Gothic", "Yu Mincho"  # Windows / Mac の PowerPoint にある和文書体
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # 表スタイル「スタイルなし、表のグリッドなし」
ALIGN = {"left": PP_ALIGN.LEFT, "start": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


class Scale:
    def __init__(self, w_px):
        self.k = SLIDE_W / w_px  # EMU / px
        self.pt = 960.0 / w_px   # pt / px（13.333in × 72pt）

    def __call__(self, v):
        return Emu(int(round(v * self.k)))


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def font_name(run):
    f = run.get("font", "")
    if run.get("serif") or re.search(r"Mincho|明朝|Serif", f, re.I):
        return SERIF
    return SANS


def strip_style(shape):
    """add_shape が付ける既定スタイル（影・テーマ色の枠）を外す"""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def fill_runs(tf, paras, sc, first=True):
    for pi, p in enumerate(paras):
        para = tf.paragraphs[0] if (first and pi == 0) else tf.add_paragraph()
        para.alignment = ALIGN.get(p.get("align", "left"), PP_ALIGN.LEFT)
        lh = p.get("lineHeight")
        if lh:
            para.line_spacing = Pt(lh * sc.pt)
        para.space_before = para.space_after = Pt(0)
        bu = p.get("bullet")
        # HTML 側で「• 」を文字で打っている行も、PowerPoint では箇条書き書式にする（slide-rules §7.3）
        if not bu and p["runs"]:
            m = re.match(r"^([•・·–])\s+", p["runs"][0]["text"])
            if m:
                p = dict(p, runs=[dict(p["runs"][0], text=p["runs"][0]["text"][m.end():])] + p["runs"][1:])
                bu = {"char": m.group(1), "color": p["runs"][0].get("color"), "size": p["runs"][0]["size"]}
        if bu:  # 行頭記号は文字で打たず箇条書き書式に（slide-rules §7.3）。ぶら下げで2行目以降を本文の頭に揃える
            first_size = p["runs"][0]["size"] if p["runs"] else 12
            hang = p.get("hang") or first_size * 1.0
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(int(sc(hang))))
            pPr.set("indent", str(-int(sc(hang))))
            for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buAutoNum", "a:buNone"):
                for el in pPr.findall(qn(tag)):
                    pPr.remove(el)
            if bu.get("color"):
                sf = etree.SubElement(pPr, qn("a:buClr"))
                etree.SubElement(sf, qn("a:srgbClr")).set("val", bu["color"].lstrip("#").upper())
            if bu.get("size") and first_size:
                etree.SubElement(pPr, qn("a:buSzPct")).set("val", str(int(min(400, max(25, bu["size"] / first_size * 100)) * 1000)))
            if bu.get("auto"):
                an = etree.SubElement(pPr, qn("a:buAutoNum"))
                an.set("type", bu["auto"])
                if bu.get("start", 1) > 1:
                    an.set("startAt", str(bu["start"]))
            else:
                etree.SubElement(pPr, qn("a:buFont")).set("typeface", SANS)
                etree.SubElement(pPr, qn("a:buChar")).set("char", bu["char"])
        for r in p["runs"]:
            run = para.add_run()
            run.text = r["text"]
            f = run.font
            f.size = Pt(round(r["size"] * sc.pt * 2) / 2)
            f.bold, f.italic = r.get("bold", False), r.get("italic", False)
            if r.get("underline"):
                f.underline = True
            f.color.rgb = rgb(r.get("color", "#000000"))
            name = font_name(r)
            f.name = name
            rPr = run._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rPr, qn(tag))
                el.set("typeface", name)


def add_rect(slide, it, sc):
    x, y, w, h = sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"])
    if w <= 0 or h <= 0:
        return
    if it.get("poly"):
        pts = [(sc(it["x"] + px), sc(it["y"] + py)) for px, py in it["poly"]]
        fb = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
        fb.add_line_segments(pts[1:], close=True)
        shp = fb.convert_to_shape()
    elif it.get("ellipse") or (it.get("radius", 0) >= min(it["w"], it["h"]) / 2 - 0.5 and abs(it["w"] - it["h"]) <= 1):
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)  # border-radius 50% → 楕円・円
    elif it.get("radius", 0) > 0.5:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shp.adjustments[0] = min(0.5, it["radius"] / max(1, min(it["w"], it["h"])))
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    strip_style(shp)
    if it.get("fill"):
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(it["fill"])
        a = it.get("alpha", 1)
        if a < 0.99:
            clr = shp.fill._xPr.find(".//" + qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(a * 100000)))
    else:
        shp.fill.background()
    ln = it.get("line")
    if ln:
        shp.line.color.rgb = rgb(ln["color"])
        shp.line.width = Pt(max(0.25, ln["w"] * sc.pt))
        if ln.get("dash"):
            from pptx.enum.dml import MSO_LINE
            shp.line.dash_style = MSO_LINE.DASH if ln["dash"] == "dashed" else MSO_LINE.ROUND_DOT
    else:
        shp.line.fill.background()
    shp.text_frame.text = ""


def add_line(slide, it, sc):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sc(it["x1"]), sc(it["y1"]), sc(it["x2"]), sc(it["y2"]))
    strip_style(c)
    c.line.color.rgb = rgb(it["color"])
    c.line.width = Pt(max(0.25, it["w"] * sc.pt))
    if it.get("dash"):
        from pptx.enum.dml import MSO_LINE
        c.line.dash_style = MSO_LINE.DASH if it["dash"] == "dashed" else MSO_LINE.ROUND_DOT


def add_text(slide, it, sc):
    # 改行がすべて明示（<br>・段落）で、自動の折り返しが無い文字は、PowerPoint でも折り返さない
    single = it.get("lines", 1) <= len(it["paras"])
    # 1行の文字は折り返さない（PowerPoint の字幅差で改行が増えるのを防ぐ）。複数行は元の折り返し幅に揃える
    # 複数行も PowerPoint の字幅差で行末の1字（句点など）が次行に落ちないよう、わずかに広げる
    w = it["w"] * (1.04 if single else 1.015) + 2
    x = it["x"]
    align = it["paras"][0].get("align", "left") if it["paras"] else "left"
    if single and align in ("center",):
        x -= (w - it["w"]) / 2
    elif single and align in ("right", "end"):
        x -= (w - it["w"])
    tb = slide.shapes.add_textbox(sc(x), sc(it["y"]), sc(w), sc(max(it["h"], 1)))
    if it.get("role") == "title":
        tb.name = "Title"  # check_deck.py がタイトルとして読む
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = not single
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    if it.get("vert"):
        tf._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
    fill_runs(tf, it["paras"], sc)


def cell_border(cell, side, b, sc):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"t": "a:lnT", "r": "a:lnR", "b": "a:lnB", "l": "a:lnL"}[side]
    old = tcPr.find(qn(tag))
    if old is not None:
        tcPr.remove(old)
    ln = etree.Element(qn(tag))
    if b:
        ln.set("w", str(int(max(0.25, b["w"] * sc.pt) * 12700)))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", b["color"].lstrip("#").upper())
        if b.get("dash"):
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash" if b["dash"] == "dashed" else "sysDot")
    else:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    # a:lnL, lnR, lnT, lnB は塗りより前・この順（スキーマ順）
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    idx = 0
    for child in list(tcPr):
        if child.tag in [qn(o) for o in order[:order.index(tag)]]:
            idx = list(tcPr).index(child) + 1
    tcPr.insert(idx, ln)


def add_table(slide, it, sc):
    nr, nc = len(it["rows"]), len(it["cols"])
    if not nr or not nc:
        return
    gf = slide.shapes.add_table(nr, nc, sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"]))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(attr, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE_TABLE
    for i, w in enumerate(it["cols"]):
        tbl.columns[i].width = sc(w)
    for i, h in enumerate(it["rows"]):
        tbl.rows[i].height = sc(h)
    for row in tbl.rows:  # 結合で隠れるセルも含め、既定の18ptを外す
        for cell in row.cells:
            cell.text_frame.paragraphs[0]._p.get_or_add_endParaRPr().set("sz", "100")
    covered = set()
    for c in it["cells"]:
        r0, c0, rs, cs = c["r"], c["c"], c["rs"], c["cs"]
        if r0 >= nr or c0 >= nc:
            continue
        cell = tbl.cell(r0, c0)
        if rs > 1 or cs > 1:
            cell.merge(tbl.cell(min(nr, r0 + rs) - 1, min(nc, c0 + cs) - 1))
        covered.add((r0, c0))
        pt, pr, pb, pl = c["pad"]
        cell.margin_top, cell.margin_right, cell.margin_bottom, cell.margin_left = sc(pt), sc(pr), sc(pb), sc(pl)
        cell.vertical_anchor = VALIGN.get(c.get("valign"), MSO_ANCHOR.MIDDLE)
        if c.get("fill"):
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(c["fill"])
        else:
            cell.fill.background()
        for side, b in zip("trbl", c["borders"]):
            cell_border(cell, side, b, sc)
        if c.get("vert"):
            cell._tc.get_or_add_tcPr().set("vert", "eaVert")
        tf = cell.text_frame
        tf.word_wrap = True
        if c["paras"]:
            fill_runs(tf, c["paras"], sc)
        else:
            # 空セル（図形を重ねるセルを含む）は既定の18ptのままだと行が伸びるので、行の高さに収まる大きさにする
            h_pt = it["rows"][r0] * sc.pt - (pt + pb) * sc.pt
            size = max(1, min(8, h_pt * 0.7))
            endp = tf.paragraphs[0]._p.get_or_add_endParaRPr()
            endp.set("sz", str(int(size * 100)))


def convert(html, out):
    here = Path(__file__).resolve().parent
    with tempfile.TemporaryDirectory() as td:
        r = subprocess.run(["node", str(here / "html_dump.mjs"), str(Path(html).resolve()), td],
                           capture_output=True, text=True, timeout=600)
        if r.returncode != 0:
            sys.exit(f"html_dump.mjs 失敗: {r.stderr.strip()[:400]}")
        data = json.loads((Path(td) / "dump.json").read_text())
        sc = Scale(data["slideW"])
        prs = Presentation()
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
        blank = prs.slide_layouts[6]
        for s in data["slides"]:
            slide = prs.slides.add_slide(blank)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = rgb(s["bg"])
            items = s["items"]
            # 重なり順: 塗り・線・画像・表は HTML の順、文字はいちばん上
            for it in items:
                t = it["t"]
                if t == "rect":
                    add_rect(slide, it, sc)
                elif t == "line":
                    add_line(slide, it, sc)
                elif t == "img":
                    p = Path(td) / it["file"]
                    if p.exists():
                        slide.shapes.add_picture(str(p), sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"]))
                elif t == "table":
                    add_table(slide, it, sc)
            for it in items:
                if it["t"] == "text" and it["paras"]:
                    add_text(slide, it, sc)
        prs.save(out)
    return len(data["slides"])


def main():
    args = sys.argv[1:]
    if not args:
        sys.exit(__doc__)
    html = args[0]
    out = args[1] if len(args) > 1 else str(Path(html).with_suffix(".pptx"))
    n = convert(html, out)
    print(f"{n} slides → {out}")
    print("次: check_deck.py で検査し、PowerPoint で開いて折り返しと重なりを確認する（slide-rules §8.6）")


if __name__ == "__main__":
    main()
