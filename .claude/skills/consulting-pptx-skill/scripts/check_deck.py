#!/usr/bin/env python3
"""スライド規約の機械チェック（PPTX / HTML 共通）。

使い方:
  python3 check_deck.py out.pptx
  python3 check_deck.py deck.html
  python3 check_deck.py assets/SuperTemplate_36type.pptx --template   # テンプレ集そのものを検査するとき
  python3 check_deck.py deck.html --forbid ~/.config/deck-forbidden-terms.txt   # 顧客名・社内語の残りを FAIL
  python3 check_deck.py pages.pptx --house house.skin.json   # 既存の資料へ差し込むページ（slide-rules §8.7）
  python3 check_deck.py out.pptx --xml-only                  # PowerPoint が開けなくなる不正だけを見る

正典: references/slide-rules.md
終了コード: FAIL があれば 1。
"""
import re
import sys
import zipfile
from pathlib import Path

# 自ブランドで禁止するレガシー色があればここに列挙（HEX 6桁・#なし）
OLD_COLORS = []
TITLE_MAX = 40
EMU_W, EMU_H = 12192000, 6858000

fails, warns = [], []


def fail(msg):
    fails.append(msg)


def warn(msg):
    warns.append(msg)


# --- PowerPoint が開けなくなる不正（slide-rules §8.7）-------------------------------
# テーマ色の欄に HEX や存在しない名前を書いた PPTX は、python-pptx では読めても PowerPoint は「修復」を求める。
# 確認のために自動で開かせると、そのダイアログが残って以後の操作を止める。開く前にここで止める。
SCHEME_COLORS = {"bg1", "tx1", "bg2", "tx2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
                 "hlink", "folHlink", "dk1", "lt1", "dk2", "lt2", "phClr"}
XML_ONLY = False   # --xml-only: 上の不正だけを見て終える（PowerPoint に開かせる前に使う）
HOUSE = None       # --house: measure_deck.py が書き出した、差し込む先の資料の書式


def check_xml_validity(idx, xml):
    bad = sorted(set(re.findall(r'<a:schemeClr val="([^"]*)"', xml)) - SCHEME_COLORS)
    if bad:
        fail(f"p{idx}: テーマ色の名前が不正 {bad}（HEX は srgbClr に書く。このままでは PowerPoint が開けない — §8.7）")
    if re.search(r'<a:ext cx="-\d+"|<a:ext cx="\d+" cy="-\d+"', xml):
        fail(f"p{idx}: 図形の大きさが負の値（PowerPoint が開けない — §8.7）")
    szs = [int(v) for v in re.findall(r'<a:(?:rPr|endParaRPr|defRPr)[^>]*\bsz="(\d+)"', xml)]
    if any(not 100 <= v <= 400000 for v in szs):
        fail(f"p{idx}: 文字の大きさ sz が範囲外（100〜400000。PowerPoint が開けない — §8.7）")


def check_house(idx, slide, xml):
    """差し込む先の資料と作りが合っているか（slide-rules §8.7）。HOUSE は measure_deck.py の出力"""
    if HOUSE.get("title") and slide.shapes.title is None:
        fail(f"p{idx}: タイトルがプレースホルダーに入っていない（位置と書体は資料のレイアウトが決める — §8.7）")
    faces = sorted(set(re.findall(r'<a:(?:latin|ea) typeface="((?!\+)[^"]*)"', xml)))
    if faces and HOUSE["fonts"]["inherit"]:
        fail(f"p{idx}: 書体を run に直指定 {faces[:4]}（資料はテーマ任せ。差し込むと書体がずれる — §8.7）")
    elif faces and set(faces) - {HOUSE["fonts"].get("explicit")}:
        warn(f"p{idx}: 資料と違う書体 {sorted(set(faces) - {HOUSE['fonts'].get('explicit')})[:4]}（資料は {HOUSE['fonts'].get('explicit')}）")
    if "<a:tbl>" in xml and not HOUSE["tables"]["native"]:
        warn(f"p{idx}: PowerPoint の表オブジェクト（資料の表はテキストボックス＋罫線。deck_pptx の table で組む — §8.7）")
    if 'type="slidenum"' in xml and HOUSE.get("page_number") == "layout":
        warn(f"p{idx}: ページ番号をスライドに置いている（資料はレイアウトが出すので二重になる — §8.7）")
    used = set(v.upper() for v in re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)) | set(re.findall(r'<a:schemeClr val="([^"]+)"', xml))
    ng = (HOUSE.get("colors") or {}).get("ng") or "FF0000"
    extra = sorted(used - set(HOUSE.get("palette") or used) - {ng.split("|")[0].upper(), "FF0000", "FFFFFF", "000000", "bg1", "tx1"})
    if extra:
        warn(f"p{idx}: 資料に無い色 {extra[:6]}（資料のテーマ色から選ぶ — §8.7）")


# テンプレ集そのものを検査するときだけ True（--template）。
# 通常のデッキでは「Text N」「◯◯」等のプレースホルダーが残っていたら FAIL にする。
TEMPLATE_MODE = False
STRICT_LEN = True  # 互換のため残置。PPTX/HTMLとも2行許容（>80字のみFAIL）

# 表記ゆれ代表ペア（slide-rules §7.6: 1資料1用語）。両方の表記が同一資料に現れたら WARN。
# (ラベルA, パターンA, ラベルB, パターンB)
PART_TYPE_NAMES = set()

TERM_VARIANTS = [
    ("メモリー", r"メモリー", "メモリ", r"メモリ(?!ー)(?!・CPU)"),  # 計算機の「メモリ・CPU」は除外
    ("紐付", r"紐付", "紐づ", r"紐づ"),
    ("ユーザー", r"ユーザー", "利用者", r"利用者"),
    ("アセット", r"アセット", "資産", r"資産"),
    ("フォルダー", r"フォルダー", "フォルダ", r"フォルダ(?!ー)"),
    ("サーバー", r"サーバー", "サーバ", r"サーバ(?!ー)"),
    ("メンバー", r"メンバー", "メンバ", r"メンバ(?!ー)"),
    ("コンピューター", r"コンピューター", "コンピュータ", r"コンピュータ(?!ー)"),
    ("問い合わせ", r"問い合わせ", "問合せ", r"問合せ"),
]


# AI臭ワード（slide-rules §7.9 / references/ai-smell-lexicon.md）。高確度語のみ WARN。
AI_SMELL_WORDS = [
    "まさに", "非常に", "極めて", "圧倒的", "画期的", "革新的", "次世代の",
    "過言ではありません", "に他なりません", "シームレス", "シナジー", "ソリューション",
    "エンドツーエンド", "ブラッシュアップ", "付加価値",
    "寄り添い", "伴走し", "二人三脚", "さらなる高みへ", "邁進",
    "昨今", "変化の激しい", "という点において", "の観点から",
    "させていただきます", "いただけますと幸いです",
    "と言えるでしょう", "と考えられます", "することが可能です",
    # 先送りの決まり文句（中身を書かずに後へ回す）
    "詳細は別途", "別途ご説明", "追ってご連絡", "今後検討してまいります", "詳細は後日",
]


def check_production_meta(slides):
    """制作メタ（本スキル名・リポジトリ）がスライドに見えていないか（SKILL.md「本スキル使用の注釈」）。
    ツール名を置けるのは最終ページ（裏表紙）の出典行だけ。それ以外のページに出ていたら FAIL。
    「本資料は…で作成」の文は免責文のこともあるので、最終ページ以外にあれば WARN で目視に回す。"""
    def visible(s):
        s = re.sub(r"<(script|style)[^>]*>.*?</\1>|<!--.*?-->", " ", s, flags=re.S)
        return re.sub(r"<[^>]+>", " ", s)
    if TEMPLATE_MODE:   # パーツ集は見本の並びなので裏表紙が最終ページにない
        return
    for i, s in enumerate(slides, 1):
        if i == len(slides):
            continue
        v = visible(s)
        m = re.search(r"consulting-pptx-skill", v)
        if m:
            fail(f"p{i}: 制作メタ情報がスライドに表示されている: 「{m.group(0)}」（ツール名を置けるのは最終ページの出典行だけ）")
            continue
        m = re.search(r"本資料は[^。<]{0,40}で作成", v)
        if m:
            warn(f"p{i}: 「{m.group(0)}」（制作クレジットなら最終ページの出典行へ。免責文なら可）")
        elif re.search(r"github\.com", v):
            warn(f"p{i}: スライドに github.com が表示されている（制作メタなら消す。出典として公開リポジトリを示すなら可）")


def pptx_text(xml):
    """スライド XML の表示テキストを段落（<a:p>）ごとに 1 行にして返す。
    run（<a:t>）ごとに改行すると「ラベル」「 — 」「説明」のように run が分かれたダッシュ連結を見落とす。"""
    paras = re.findall(r"<a:p\b.*?</a:p>", xml, re.S)
    if not paras:
        return "\n".join(re.findall(r"<a:t>(.*?)</a:t>", xml, re.S))
    return "\n".join("".join(re.findall(r"<a:t>(.*?)</a:t>", p, re.S)) for p in paras)


def check_ai_smell(pages):
    """slide-rules §7.9: AI臭の高確度語を検出（WARN。文脈上正当なら目視で無視してよい）"""
    hits = {}
    for i, txt in pages:
        found = [w for w in AI_SMELL_WORDS if w in txt]
        if found:
            hits[i] = found
    if hits:
        detail = "、".join(f"p{i}:「{'/'.join(ws[:3])}」" for i, ws in sorted(hits.items())[:6])
        warn(f"AI臭ワード検出: {detail}（§7.9 / ai-smell-lexicon.md。素の動詞・直球の言い方に置き換え）")
    # 非該当セルの「—」単独（§6）は連結ではないので除外する。行（文章塊）単位で、
    # 「—」が他の文字と同じ塊に入っているものだけを数える（「ラベル — 説明」「A—B」）。
    def _dash_joined(txt):
        for line in txt.split("\n"):
            t = line.strip()
            if "—" not in t or re.fullmatch(r"—+(?:\s*[（(][^）)]*[）)])?", t):
                continue
            return True
        return False
    dash_pages = sorted({i for i, txt in pages if _dash_joined(txt)})
    if dash_pages:
        warn(f"ダッシュ「 — 」連結 p{dash_pages}（AI文体の典型。句点・「：」・括弧に置き換え。§7.9）")


def check_terms(pages):
    """pages: [(idx, text), ...] 資料全体で両方の表記が出たら WARN（意味が別なら目視で無視してよい）"""
    for la, pa, lb, pb in TERM_VARIANTS:
        hits_a = sorted({i for i, t in pages if re.search(pa, t)})
        hits_b = sorted({i for i, t in pages if re.search(pb, t)})
        if hits_a and hits_b:
            warn(f"表記ゆれ疑い: 「{la}」p{hits_a} と「{lb}」p{hits_b} が混在（§7.6 1資料1用語。別概念なら可・目視確認）")

# --- 数値の平仄（slide-rules §7.6）--------------------------------------------------
# 同じ指標（数の直前の語）が、別のページで違う値で書かれていたら WARN。
# 例:「売上高 120億円」(p3) と「売上高は約118億円」(p7)、「利用者数 4.2万人」と「利用者数 42,500人」。
# 単位の桁（千・万・億・兆）は掛けて比べるので、書き方が違うだけで値が同じなら出さない。
# 時点（2024年／2024年度）が直前にあれば指標名に含めて比べるので、年の違う同じ指標は食い違いにしない。
NUM_UNITS = (r"円|ドル|ユーロ|人|名|件|社|団体|店舗|拠点|校|戸|世帯|台|個|本|枚|冊|回|倍|%|％|pt|ポイント|"
             r"時間|日|か月|ヶ月|週|kWh|MWh|GWh|kW|MW|GW|kg|km|m3|㎥|m2|㎡|ha|ヘクタール|g|t|トン|m|L")
NUM_LABEL = re.compile(
    r"(?:(?P<year>(?:19|20)\d{2})年度?(?:の|時点の|末の|末時点の)?)?"
    r"(?P<label>[一-龥々ァ-ヶー]{2,12})(?:は|が|の|：|:|＝|=)?\s*(?:約|およそ|計|合計)?\s*"
    r"(?P<num>\d[\d,]*(?:\.\d+)?)\s*(?P<unit>万|億|兆|千)?(?P<base>" + NUM_UNITS + r")(?![A-Za-z])")
_ZEN = str.maketrans("０１２３４５６７８９，．％：＝", "0123456789,.%:=")
_MULT = {"千": 1e3, "万": 1e4, "億": 1e8, "兆": 1e12}
_UNIT_ALIAS = {"名": "人", "％": "%", "ポイント": "pt", "ヶ月": "か月", "㎥": "m3", "㎡": "m2",
               "ヘクタール": "ha", "トン": "t"}
# 指標名にならない一般語（「以上 80%」「平均2.1%」など）は比べない
_GENERIC_LABELS = {"以上", "以下", "未満", "超", "平均", "合計", "全体", "最大", "最小", "最高", "最低",
                   "前年", "前年比", "同期", "うち", "残り", "目標", "実績", "約", "計"}


def _num_facts(text):
    """(指標名, 単位) -> 値 の組を拾う。値は単位の桁（万・億）を掛けた数で比べる。"""
    out = []
    for m in NUM_LABEL.finditer(text.translate(_ZEN)):
        if m["label"] in _GENERIC_LABELS:
            continue
        label = (m["year"] + "年:" if m["year"] else "") + m["label"]
        base = _UNIT_ALIAS.get(m["base"], m["base"])
        try:
            val = round(float(m["num"].replace(",", "")) * _MULT.get(m["unit"] or "", 1), 6)   # 1.1億 = 110,000,000 を float 誤差で食い違いにしない
        except ValueError:
            continue
        out.append(((label, base), val, m.group(0).strip()))
    return out


def check_number_consistency(pages):
    """pages: [(idx, text), ...] 同じ指標がページ間で違う値なら WARN（別概念・別時点なら目視で無視してよい）"""
    seen = {}
    for i, t in pages:
        for key, val, raw in _num_facts(t):
            seen.setdefault(key, []).append((i, val, raw))
    for (label, base), hits in seen.items():
        vals = {v for _, v, _ in hits}
        pages_hit = {i for i, _, _ in hits}
        if len(vals) > 1 and len(pages_hit) > 1:
            detail = " / ".join(f"p{i}「{raw}」" for i, _, raw in hits[:4])
            warn(f"数値の平仄疑い: 「{label.replace(':', '')}」が {detail} で食い違う（§7.6 数値の平仄。別時点・別範囲なら注記して可）")

# --- 本文のプレースホルダー残り（slide-rules §2.8 / README）-------------------------
# タイトルだけでなく本文・表・カードに「Text N」「ラベル N」「YYYY」「パーツNN｜」が残っていたら FAIL。
# 型名をそのままタイトルにしたページ（例:「軸のある表」）も FAIL。
BODY_PLACEHOLDER = re.compile(
    r"Text\s*\d+|ラベル\s*\d+|タイトル\s*\d+|Source\s*\d+|YYYY|パーツ\s*\d+\s*[｜|]|ダミー|^会社名$|^連絡先$")


def _part_type_names():
    """パーツ集の型名（new_deck.py --list と同じ並び）。テンプレが無い環境では空。"""
    names = set()
    root = Path(__file__).resolve().parent.parent / "templates"
    for f in ("freeform_parts_16x9.html", "freeform_parts_more_16x9.html"):
        fp = root / f
        if not fp.exists():
            continue
        for m in re.finditer(r"パーツ\s*\d+\s*[｜|]\s*([^<]+)<", fp.read_text(encoding="utf8", errors="ignore")):
            names.add(m.group(1).strip())
    return names


def check_body_placeholders(idx, leaf_texts, title):
    if TEMPLATE_MODE:
        return
    hits = sorted({t for t in leaf_texts if BODY_PLACEHOLDER.search(t)})
    if hits:
        fail(f"p{idx}: 本文にテンプレのプレースホルダーが残っている ×{len(hits)}: {' / '.join(h[:20] for h in hits[:4])}")
    if title and title.strip() in PART_TYPE_NAMES:
        fail(f"p{idx}: タイトルがパーツの型名のまま「{title.strip()}」（主張文に書き換える — §2.8）")


# --- 1ブロック1文（文章を連続して詰め込まない）------------------------------------
# スライドの1つの文章塊（セル・カード本文・段落・箇条1行）に2文以上を入れない。
# 2文以上になる中身は1項目1文の箇条に分ける。機械判定は「句点で終わる文が2つ以上」。
SENTENCE_END = re.compile(r"[。！？!?](?=\s*\S)")


def check_multi_sentence(idx, leaf_texts):
    bad = []
    for t in leaf_texts:
        if t.startswith(("出典", "注", "※", "Source")):
            continue  # 出典・注記行は対象外
        core = re.sub(r"（[^）]*）|\([^)]*\)|「[^」]*」", "", t)  # 括弧・引用内の句点は数えない
        if len(SENTENCE_END.findall(core.strip())) >= 1:
            bad.append(t)
    if bad:
        fail(f"p{idx}: 1つの文章塊に2文以上 ×{len(bad)}: 「{bad[0][:36]}…」（1項目1文の箇条に分ける。動きはきっかけ→起きること→判断→結果の順）")


# --- 禁止語（顧客名・社内語・案件コード）-------------------------------------------
# 公開・社外共有前に、資料ごとに出してはいけない語を外部ファイルで渡す（リポジトリには入れない）。
#   python3 check_deck.py deck.html --forbid ~/.config/deck-forbidden-terms.txt
# 1行1語。# で始まる行はコメント。re: で始めると正規表現。
FORBIDDEN_TERMS = []


def load_forbidden(path):
    terms = []
    for line in Path(path).expanduser().read_text(encoding="utf8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        terms.append(re.compile(line[3:]) if line.startswith("re:") else re.compile(re.escape(line)))
    return terms


def check_forbidden(pages, raw=""):
    if not FORBIDDEN_TERMS:
        return
    for pat in FORBIDDEN_TERMS:
        hit_pages = sorted({i for i, t in pages if pat.search(t)})
        in_markup = bool(raw and pat.search(re.sub(r">[^<]*<", "><", raw)))  # コメント・属性・alt 等
        if hit_pages or in_markup:
            where = f"p{hit_pages}" if hit_pages else "HTMLのコメント/属性"
            # 語そのものは出力に出さない（ログ経由の再流出を防ぐ）
            fail(f"禁止語リストの語が {where} に残っている（リスト{FORBIDDEN_TERMS.index(pat)+1}行目の語）")


def check_title(idx, title, explicit_break=False):
    # explicit_break: 意味の切れ目で明示改行済み（§2.13）なら2行想定の WARN は出さない
    t = title.strip()
    if not t:
        warn(f"p{idx}: タイトルが空（表紙/扉なら可）")
        return
    if re.search(r"(です|ます|でした|ました)[。．.]?$", t):
        fail(f"p{idx}: タイトルがですます調 → 体言止めに: 「{t}」")
    # タイトルは PPTX/HTML とも2行まで許容。1行目安(TITLE_MAX=40字)超は WARN、2行にも収まらない長さ(>80字)だけ FAIL。
    # 文字を縮小して1行に詰めるのは不可（slide-rules §2.1）。
    # 半角文字（英数・記号・空白）は全角の半分として数える（英語タイトルを日本語基準で FAIL にしない）
    tlen = sum(0.5 if ord(ch) < 0x3000 else 1 for ch in t)
    if tlen > TITLE_MAX * 2:
        fail(f"p{idx}: タイトル 全角換算{tlen:.0f}字（>{TITLE_MAX*2}・2行にも収まらない。主張を絞る）: 「{t}」")
    elif tlen > TITLE_MAX and not explicit_break:
        warn(f"p{idx}: タイトル 全角換算{tlen:.0f}字（2行になる想定。意味の切れ目で改行・泣き別れなし・文字縮小で1行に詰めない）: 「{t}」")
    if re.match(r"^(Step|STEP|ステップ)\s*\d", t):
        fail(f"p{idx}: タイトルに Step 連結（タグチップで表現）: 「{t}」")
    if re.search(r"^(この|その|ここまで)", t):
        fail(f"p{idx}: 他スライド参照語で始まるタイトル: 「{t}」")
    if re.match(r"^(まずは|では|そして|さらに|ちなみに)|^(まず|また|次に)[、,]", t):
        warn(f"p{idx}: タイトルが話し言葉の接続詞で始まる（主語から書く — slide-rules §2.18）: 「{t}」")
    if t.count("（") + t.count("(") >= 2:
        warn(f"p{idx}: タイトルに丸括弧が多い: 「{t}」")
    if not TEMPLATE_MODE and re.search(r"[◯○]{2,}|Text\s*\d|ラベル\s*\d|タイトル\s*\d|Source\s*\d|YYYY|ダミー|^資料名$|^会社名$", t):
        fail(f"p{idx}: テンプレのプレースホルダーが残っている（タイトルはストーリーラインから書く — slide-rules §2.8）: 「{t}」")


# テンプレの見本タイトルは「形の参考」であって埋める鋳型ではない（slide-rules §2.8）。同じ文型が並んだら WARN。
def check_title_variety(titles):
    real = [t.strip() for t in titles if t and t.strip()]
    if len(real) < 5:
        return
    molds = {
        "「◯◯は、…」": lambda t: re.match(r"^.{1,12}は[、,]", t),
        "「◯◯には、…がある」": lambda t: re.search(r"には[、,].*(ある|存在する)$", t),
        "数を主語にした形（「3つの…」）": lambda t: re.match(r"^[0-9０-９一二三四五六七八九十]+つ", t),
    }
    for name, f in molds.items():
        hits = [t for t in real if f(t)]
        if len(hits) >= max(4, int(len(real) * 0.6)):
            warn(f"タイトルの文型が {name} に偏っている（{len(hits)}/{len(real)}枚）。テンプレの見本文型をなぞらず、主張ごとに自然な文で書く — slide-rules §2.8")


# --- 数の不一致（slide-rules §2.9） -------------------------------------------
# タイトルに書いた数と、本文の連番ラベルの最大値を突き合わせる。
# 例: タイトル「3段階で移行する」に対し本文が「STEP 1〜4」。
# 機械で拾える数少ない「内容の矛盾」なので FAIL にする。
KANSUJI = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
           "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
COUNT_WORDS = r"(?:段階|フェーズ|ステップ|柱|論点)"
ENUM_FAMILY = [
    (r"段階|フェーズ|ステップ", r"(?:フェーズ|ステップ|STEP|Step|PHASE|Phase|段階)"),
    (r"柱", r"(?:柱)"),
    (r"論点", r"(?:論点)"),
]


def _to_int(s):
    s = s.translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    if s.isdigit():
        return int(s)
    return KANSUJI.get(s)


def check_count_match(idx, title, text):
    """タイトルの「N段階」等が本文の連番ラベルと食い違っていたら FAIL"""
    if not title or not text:
        return
    for m in re.finditer(r"([0-9０-９]+|[一二三四五六七八九十])\s*(" + COUNT_WORDS + ")", title):
        n = _to_int(m.group(1))
        if not n or not (2 <= n <= 12):
            continue
        fam = next((lab for pat, lab in ENUM_FAMILY if re.search(pat, m.group(2))), None)
        if not fam:
            continue
        found = [v for v in (_to_int(x.group(1))
                             for x in re.finditer(fam + r"\s*([0-9０-９]+)", text)) if v]
        if found and max(found) != n:
            fail(f"p{idx}: タイトルの「{m.group(0)}」と本文の連番（最大{max(found)}）が食い違う"
                 f"（§2.9 数はページの中身と一致させる）: 「{title}」")


# ---------------------------------------------------------------- PPTX
# ---- スキル有無検証（slide-rules §2.13 泣き別れ／§5.13 版面充填率）
def _fwlen(t):
    return sum(0.5 if ord(ch) < 0x3000 else 1 for ch in t)

def check_orphan(idx, shape, label="タイトル"):
    """箱幅とフォントサイズから1行容量を推定し、最終行が1〜3字だけになる折返し（泣き別れ）を WARN"""
    try:
        if not (shape.has_text_frame and shape.width):
            return
        tf = shape.text_frame
        if tf.word_wrap is False:  # 折り返しなしの文字（html_to_pptx.py の1行テキスト等）は泣き別れしない
            return
        raw = tf.text
        if not raw.strip():
            return
        szs = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
        if not szs:
            return
        pt = max(szs)
        w_in = shape.width / 914400.0
        cap = max(4, int(w_in / (pt / 72.0)))
        for line in raw.replace("\v", "\n").split("\n"):
            L = _fwlen(line.strip())
            if L <= cap:
                continue
            last = L - cap * ((int(-(-L // cap))) - 1)
            if 0 < last <= 3:
                warn(f"p{idx}: {label}が泣き別れの疑い（推定{int(-(-L // cap))}行目が約{last:.0f}字だけ。意味の切れ目で明示改行 — slide-rules §2.13）: 「{line.strip()[:40]}」")
    except Exception:
        return

def check_fill_ratio(idx, slide, title_shape):
    """本文（タイトル下〜出典行上）の縦幅に対する本文要素の占有率。55%未満は WARN（§5.13）"""
    try:
        top_lim = 1.6 * 914400
        bot_lim = 6.8 * 914400
        tops, bots = [], []
        for sh in slide.shapes:
            if sh is title_shape or sh.top is None or sh.height is None:
                continue
            t, b = sh.top, sh.top + sh.height
            if b <= top_lim or t >= bot_lim:
                continue  # キッカー/タイトル/フッター
            # 幅いっぱいの罫線（フッター罫・タイトル罫）は除外
            if sh.height < 914400 * 0.02 and sh.width and sh.width > 914400 * 11:
                continue
            tops.append(max(t, top_lim)); bots.append(min(b, bot_lim))
        if not tops:
            return
        cov = (max(bots) - min(tops)) / (bot_lim - top_lim)
        if cov < 0.55:
            warn(f"p{idx}: 版面充填率 {cov*100:.0f}%（本文が版面の半分未満。情報を足すか型を変える。飾りで埋めない — slide-rules §5.13）")
    except Exception:
        return


def check_pptx(path):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("python-pptx が必要: pip3 install python-pptx")
    from pptx_open import open_presentation
    prs = open_presentation(path)
    want_w, want_h = (HOUSE["slide"]["emu"] if HOUSE else (EMU_W, EMU_H))   # 差し込むページは資料の大きさに合わせる（§1）
    if not XML_ONLY and (abs(prs.slide_width - want_w) > 2000 or abs(prs.slide_height - want_h) > 2000):
        fail(f"スライドサイズ {prs.slide_width}x{prs.slide_height} ≠ {'資料' if HOUSE else '16:9'} {want_w}x{want_h}")

    titles = []
    term_pages = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        for n in sorted(names, key=lambda s: int(re.search(r"slide(\d+)", s).group(1))):
            xml = z.read(n).decode("utf8", "ignore")
            idx = int(re.search(r"slide(\d+)", n).group(1))
            term_pages.append((idx, pptx_text(xml)))
            check_xml_validity(idx, xml)
            if re.search(r'<p:sld\b[^>]*\bshow="0"', xml):
                warn(f"p{idx}: 非表示のスライド（使わないならファイルから消す。付録の控えとして意図して残すなら無視してよい — slide-rules §8）")
            if XML_ONLY:
                continue
            # 角丸（高さ 0.4in=365760 EMU 以上の図形のみ）
            rr = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                if 'prst="roundRect"' in sp:
                    h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', sp)
                    if h and int(h.group(1)) >= 365760:
                        rr += 1
            if rr:
                fail(f"p{idx}: roundRect の大きなボックス ×{rr}（直角 rect に）")
            # 塗りありボックスに枠線（大きな図形のみ）
            fb = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                spPr = re.search(r"<p:spPr>.*?</p:spPr>", sp, re.S)
                if not spPr:
                    continue
                pr = spPr.group(0)
                h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', pr)
                if not (h and int(h.group(1)) >= 365760):
                    continue
                ln = re.search(r"<a:ln[ >].*?</a:ln>", pr, re.S)
                filled = "<a:solidFill>" in re.sub(r"<a:ln[ >].*?</a:ln>", "", pr, flags=re.S)
                if filled and ln and "<a:solidFill>" in ln.group(0):
                    fb += 1
            if fb:
                warn(f"p{idx}: 塗りあり図形に枠線 ×{fb}（塗りカードは line なし — slide-rules §5.3）")
            for c in OLD_COLORS:
                if f'val="{c}"' in xml or f'val="{c.lower()}"' in xml:
                    fail(f"p{idx}: 禁止色 {c}")
            # ブレット記号のテキスト直打ち（slide-rules §7.3: マーカーは buChar 書式で付与）
            lb = len(re.findall(r"<a:t>\s*•", xml))
            if lb:
                fail(f"p{idx}: ブレット「•」を本文テキストに直打ち ×{lb}（buChar/buNone 書式に — slide-rules §7.3）")
            ld = len(re.findall(r"<a:t>\s*[–‐-]\s\s", xml))
            if ld:
                warn(f"p{idx}: 第2階層マーカー「– 」らしきテキスト直打ち ×{ld}（buChar 書式に — slide-rules §7.3）")
        # theme / master も色チェック
        for n in z.namelist():
            if "theme" in n or "slideMaster" in n or "slideLayout" in n:
                xml = z.read(n).decode("utf8", "ignore")
                for c in OLD_COLORS:
                    if f'val="{c}"' in xml:
                        warn(f"{n}: 禁止色 {c}")

    if XML_ONLY:
        return titles
    if HOUSE:
        with zipfile.ZipFile(path) as z:
            for i, s in enumerate(prs.slides, 1):
                check_house(i, s, z.read(s.part.partname.lstrip("/")).decode("utf8", "ignore"))
    for i, s in enumerate(prs.slides, 1):
        title = ""
        named = next((sh for sh in s.shapes if sh.has_text_frame and sh.name.startswith("Title")), None)
        if s.shapes.title is not None and s.shapes.title.has_text_frame:
            title = s.shapes.title.text_frame.text
        elif named is not None:  # html_to_pptx.py が「Title」と名付けたテキストボックス
            title = named.text_frame.text
        else:
            # タイトルPH が無いビルダー: 上部 y<1.2in の最大フォントテキストをタイトルとみなす
            cands = []
            for sh in s.shapes:
                if sh.has_text_frame and sh.top is not None and sh.top < Emu(1097280):
                    sz = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
                    cands.append((sz, sh.text_frame.text))
            if cands:
                title = max(cands)[1]
        # 泣き別れ推定（タイトル＋表紙の大きな文字）と版面充填率
        title_shape = None
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text == title:
                title_shape = sh
                break
        if title_shape is not None:
            check_orphan(i, title_shape, "タイトル")
        if i == 1:
            for sh in s.shapes:
                if sh is not title_shape and sh.has_text_frame:
                    szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    if szs and max(szs) >= 13 and len(sh.text_frame.text) < 80:
                        check_orphan(i, sh, "表紙の文字")
        else:
            check_fill_ratio(i, s, title_shape)
        explicit_break = "\n" in title.strip() and all(_fwlen(l.strip()) <= TITLE_MAX for l in title.split("\n"))
        title = title.replace("\n", " ")
        titles.append(title)
        check_title(i, title, explicit_break)
        check_count_match(i, title, dict(term_pages).get(i, ""))
        # サブタイトル疑い: タイトル直下 (y 1.2〜1.75in) の細字テキスト1行
        for sh in s.shapes:
            if sh.has_text_frame and sh.top is not None and Emu(1097280) <= sh.top < Emu(1600200):
                txt = sh.text_frame.text.strip()
                if txt and "\n" not in txt and len(txt) < 60 and txt != title:
                    szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    if szs and max(szs) <= 12 and not any(r.font.bold for p in sh.text_frame.paragraphs for r in p.runs):
                        warn(f"p{i}: タイトル直下にサブタイトルらしき行: 「{txt}」")
    check_terms(term_pages)
    check_number_consistency(term_pages)
    check_ai_smell(term_pages)
    return titles


# ---------------------------------------------------------------- HTML
# §7.22 英字大文字の装飾キッカー／§4.49 接続詞で始まる左右カラム見出し
def check_kicker_and_conclusion(html):
    body = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.S)
    texts = [re.sub(r"\s+", " ", t).strip() for t in re.findall(r">([^<>]{3,60})<", body)]
    kick = [t for t in texts if re.fullmatch(r"(?:\d{2}\s*[・·/|]\s*)?[A-Z][A-Z0-9 &/·・\-]{5,}", t)
            and re.search(r"[A-Z]{3,}\s+[A-Z]{2,}", t)
            and not re.search(r"(?i)confidential|appendix|section|step|page", t)]
    if kick:
        warn(f"英字大文字の装飾キッカー ×{len(kick)}: {' / '.join(sorted(set(kick))[:4])}（§7.22: 日本語デッキでは右上タグチップで話題を示す）")
    # 左右2カラムの見出し（h3/h4/.hd）だけを見る。th・行見出し（.rh）は行軸で通して読めるので対象外（§4.49）
    heads = [re.sub(r"<[^>]+>", "", t).strip() for t in re.findall(r"<(?:h3|h4|div class=\"(?:hd|colhd|colh|col-h)[^\"]*\")[^>]*>(.*?)</", body, re.S)]
    dakara = [t for t in heads if re.match(r"^(だから|なので|つまり)[、:：]?", t)]
    if dakara:
        warn(f"左右カラムの見出しが接続詞で始まる ×{len(dakara)}（§4.49: 2コンテンツの見出しは単独で読める名詞句に）")


def _leaf_texts(fragment):
    """スライド内の文章塊（ブロック要素ごとのテキスト）。<br> と箇条記号「•」は区切りとして扱う。"""
    frag = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", fragment, flags=re.S)
    frag = re.sub(r"<!--.*?-->", " ", frag, flags=re.S)
    frag = re.sub(r"<br\s*/?>", "\n", frag)
    blocks = re.split(r"</?(?:div|p|li|td|th|h[1-6]|section|ul|ol|tr|table)\b[^>]*>", frag)
    out = []
    for b in blocks:
        for piece in re.split(r"\n|•", re.sub(r"<[^>]+>", "", b)):
            piece = re.sub(r"\s+", " ", piece).strip()
            if piece:
                out.append(piece)
    return out


def check_exec_summary(idx, title, fragment):
    """§7.16 エグゼクティブサマリーはページタイトルの羅列にせず、主文＋インデントした詳細の階層ブレットで書く"""
    if not re.search(r"エグゼクティブサマリー|エグゼクティブ・サマリー|Executive Summary", title, re.I):
        return
    if not re.search(r"<li\b(?:(?!</li>).)*?<(?:ul|ol)\b", fragment, re.S):
        warn(f"p{idx}: エグゼクティブサマリーに入れ子のブレットが無い（§7.16: 各行を主文＋インデントした詳細2〜3本で書く。ページタイトルの羅列にしない）")


def check_html(path):
    global STRICT_LEN
    STRICT_LEN = False
    html = Path(path).read_text(encoding="utf8", errors="ignore")
    if not ("338.67mm" in html and "190.5mm" in html):
        fail("16:9 サイズ（338.67mm×190.5mm）が CSS に見当たらない（A4 / 297×167 は旧仕様）")
    if re.search(r"@page\s*{[^}]*297mm", html):
        fail("@page が A4 横のまま")
    for c in OLD_COLORS:
        if re.search(c, html, re.I):
            fail(f"禁止色 #{c}")
    # 角丸
    for m in re.finditer(r"([^{}]{0,80}){[^}]*?border-radius\s*:\s*(\d+(?:\.\d+)?)(px|mm|rem|em)", html):
        sel, v, u = m.group(1), float(m.group(2)), m.group(3)
        if re.search(r"pill|chip|tag|badge|dot", sel, re.I):
            continue  # 小ピルは丸可
        px = v * {"px": 1, "mm": 3.78, "rem": 16, "em": 16}[u]
        if px >= 4:
            fail(f"border-radius {v}{u} on `{sel.strip()[-40:]}`（角丸禁止。小ピル以外は直角）")
            break
    # サブタイトル・図表ラベル
    for cls in ["figttl", "subtitle"]:
        if re.search(r'class="[^"]*\b' + cls + r'\b', html):
            fail(f"サブタイトル/図表ラベル系クラス `.{cls}` が残っている")
    for cls in ["sub", "lead", "caption"]:
        n = len(re.findall(r'class="[^"]*\b' + cls + r'\b', html))
        if n:
            warn(f"`.{cls}` ×{n} — 表紙サブタイトルなら可。コンテンツスライドのタイトル直下なら禁止（目視確認）")
    if re.search(r'<span class="ac">', html):
        fail("タイトル内の色分け <span class=\"ac\"> が残っている")
    # 表ヘッダー
    th = re.search(r"\bth\s*{[^}]*font-size\s*:\s*(\d+)px", html)
    td = re.search(r"\btd\s*{[^}]*font-size\s*:\s*(\d+)px", html)
    if th and td and int(th.group(1)) < int(td.group(1)) + 2:
        fail(f"表ヘッダー {th.group(1)}px が本文 {td.group(1)}px +2pt 未満")
    elif th and td and int(th.group(1)) < int(td.group(1)) + 3:
        warn(f"表ヘッダー {th.group(1)}px は本文 {td.group(1)}px +3px 未満（§6: 見出しは本文より+3〜4pt 大きくするのが目安）")
    if re.search(r"\bth\s*{[^}]*color\s*:\s*#?(9[0-9a-f]{5}|a[0-9a-f]{5}|b[0-9a-f]{5}|c[0-9a-f]{5}|888|999|aaa|bbb|ccc|gr[ae]y)\b", html, re.I):
        warn("表ヘッダーが薄グレー（§6: 見出しは本文と同じ濃色）")
    check_kicker_and_conclusion(html)
    # ハーベイボール ¾ の描画（中心点の無い多角形は斜めに欠けた形になる）
    if re.search(r"\.q3\s+i\s*{[^}]*clip-path\s*:\s*polygon\(\s*50%\s+0\s*,", html):
        fail("ハーベイボール ¾ の clip-path に中心点（50% 50%）が無い。左上が斜めに欠けた形になる（polygon(50% 50%,50% 0,100% 0,100% 100%,0 100%,0 50%) に直す）")
    if not re.search(r"<meta[^>]+charset\s*=\s*[\"']?utf-?8", html, re.I):
        fail('<meta charset="utf-8"> が無い（Windows のブラウザで文字化けする — slide-rules §8）')
    if re.search(r"\bth\s*{[^}]*font-weight\s*:\s*(400|normal|300)", html):
        fail("表ヘッダーが細字")
    if re.search(r"tr:nth-child\((even|odd)\)", html):
        fail("ゼブラ縞が残っている")
    # 枠線ルール（slide-rules §5.3-5.4 / §6）
    if re.search(r"\btd\b[^{}]*{[^}]*border-bottom\s*:", html) and not re.search(
            r"tr:last-child[^{}]*{[^}]*border(-bottom)?\s*:\s*(0|none)", html):
        fail("最終行の罫線が消えていない（`tr:last-child td{border-bottom:0}` を追加 — 行き先のない罫線禁止）")
    body_html = html.split("</style>", 1)[1] if "</style>" in html else html
    used_classes = set(re.findall(r'class="([^"]*)"', body_html))
    used_tokens = set(tok for cl in used_classes for tok in cl.split())
    for m in re.finditer(r"([^{}]{0,80}){([^}]*)}", html):
        sel, body = m.group(1), m.group(2)
        if re.search(r"pill|chip|tag|badge|dot|legend", sel, re.I):
            continue
        # 本文で使っていないクラスの規則は対象外（パーツ集のCSSをまとめて取り込んだデッキで誤検知しない）
        cls_in_sel = re.findall(r"\.([A-Za-z0-9_-]+)", sel)
        if cls_in_sel and cls_in_sel[-1] not in used_tokens:  # 主語（末尾のクラス）が本文に無ければ対象外
            continue
        has_fill = re.search(r"background(-color)?\s*:\s*(?!none|transparent)#?\w", body)
        has_border = re.search(r"border\s*:\s*(?!0|none)\d", body)
        if has_fill and has_border and re.search(r"card|box|pillar|mem|step|stat", sel, re.I):
            warn(f"塗りありボックスに枠線: `{sel.strip()[-40:]}`（塗りカードは border:0 — slide-rules §5.3）")
    # タイトル抽出
    titles = []
    pat = r'<(?:section|div)[^>]*class="(?:[^"]*\bslide\b[^"]*|s|s [^"]*)"[^>]*>'
    parts = re.split(pat, html)
    slides = parts[1:] if len(parts) > 1 else []
    slide_classes = re.findall(pat.replace('class="(?:', 'class="((?:', 1).replace(')"[^>]*>', '))"[^>]*>', 1), html)
    for i, s in enumerate(slides, 1):
        m = re.search(r"<h1[^>]*>(.*?)</h1>", s, re.S) or re.search(r'class="[^"]*\b(?:ttl|title|msg)\b[^"]*"[^>]*>(.*?)</', s, re.S)
        t = re.sub(r"<[^>]+>", "", m.group(1)).strip() if m else ""
        t = re.sub(r"\s+", " ", t)
        titles.append(t)
        check_title(i, t)
        check_count_match(i, t, re.sub(r"<[^>]+>", " ", s))
        leaves = _leaf_texts(s)
        check_body_placeholders(i, leaves, t)
        check_exec_summary(i, t, s)
        if "cover" not in (slide_classes[i - 1] if i - 1 < len(slide_classes) else ""):
            check_multi_sentence(i, leaves)
    if not slides:
        warn("`.slide` 要素が見つからない（タイトル検査スキップ）")
    check_production_meta(slides if slides else [html])
    if slides:
        check_terms([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)])
        check_number_consistency([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)])
        check_ai_smell([(i, re.sub(r"<[^>]+>", "\n", s)) for i, s in enumerate(slides, 1)])
        check_forbidden([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)], html)
    else:
        check_terms([(1, re.sub(r"<[^>]+>", " ", html))])
        check_ai_smell([(1, re.sub(r"<[^>]+>", "\n", html))])
    return titles


def main():
    global TEMPLATE_MODE
    args = sys.argv[1:]
    global FORBIDDEN_TERMS, PART_TYPE_NAMES, XML_ONLY, HOUSE
    if "--xml-only" in args:
        XML_ONLY = True
        args.remove("--xml-only")
    if "--house" in args:
        k = args.index("--house")
        import json
        HOUSE = json.loads(Path(args[k + 1]).read_text(encoding="utf8"))
        del args[k:k + 2]
    if "--template" in args:
        TEMPLATE_MODE = True
        args.remove("--template")
    if "--forbid" in args:
        k = args.index("--forbid")
        FORBIDDEN_TERMS = load_forbidden(args[k + 1])
        del args[k:k + 2]
    PART_TYPE_NAMES = _part_type_names()
    if not args:
        sys.exit(__doc__)
    p = args[0]
    titles = check_pptx(p) if p.lower().endswith((".pptx", ".potx")) else check_html(p)
    if not TEMPLATE_MODE:
        check_title_variety(titles)
    print("=== タイトル一覧（上から通し読みしてストーリーが繋がるか確認） ===")
    for i, t in enumerate(titles, 1):
        print(f"{i:>3}  {t or '(なし)'}")
    print()
    for w in warns:
        print("WARN ", w)
    for f in fails:
        print("FAIL ", f)
    print(f"\n{len(fails)} FAIL / {len(warns)} WARN")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
